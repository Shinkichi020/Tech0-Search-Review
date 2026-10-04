"""各種ファイル（txt, pdf, docx, xlsx, pptx）から本文テキストを抽出するモジュール"""

import os
from pathlib import Path
from docx import Document
import openpyxl
from pptx import Presentation
from pypdf import PdfReader


def extract_text_from_file(file_path: str) -> str:
    """指定されたパスのファイルからテキストを読み出して返す"""
    path = Path(file_path)
    ext = path.suffix.lower()
    text = ""

    try:
        # テキストファイル (.txt)
        if ext == ".txt":
            with open(path, "rb") as f:
                raw_bytes = f.read()

            encodings = [
                "utf-8-sig",
                "utf-8",
                "cp932",
                "shift_jis",
                "utf-16",
                "euc-jp",
            ]
            for enc in encodings:
                try:
                    decoded = raw_bytes.decode(enc)
                    if decoded.strip():
                        text = decoded
                        break
                except Exception:
                    continue

            if not text and raw_bytes:
                text = raw_bytes.decode("utf-8", errors="replace")

        # PDF ファイル (.pdf)
        elif ext == ".pdf":
            reader = PdfReader(path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"

        # Word ファイル (.docx)
        elif ext == ".docx":
            doc = Document(path)
            text = "\n".join([p.text for p in doc.paragraphs if p.text])

        # Excel ファイル (.xlsx) — 全セルを確実に抽出する強化版
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(path, data_only=True)
            excel_texts = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                sheet_text = f"【シート名: {sheet}】\n"
                rows_data = []
                for row in ws.iter_rows(values_only=True):
                    # 空セルを除外して文字列化
                    row_cells = [
                        str(cell).strip()
                        for cell in row
                        if cell is not None and str(cell).strip() != ""
                    ]
                    if row_cells:
                        rows_data.append(" | ".join(row_cells))

                if rows_data:
                    sheet_text += "\n".join(rows_data) + "\n"
                    excel_texts.append(sheet_text)

            text = "\n\n".join(excel_texts)

        # PowerPoint ファイル (.pptx)
        elif ext == ".pptx":
            prs = Presentation(path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text += shape.text + "\n"

    except Exception as e:
        print(f"Error reading {file_path}: {e}")

    return text.strip()


def load_all_documents(folder_path: str = "dummy_data") -> list[dict]:
    """指定フォルダ内の全サポートファイルを一括読み込みしてリストで返す"""
    documents = []
    folder = Path(folder_path)

    if not folder.exists():
        return documents

    supported_exts = {".txt", ".pdf", ".docx", ".xlsx", ".pptx"}

    for file_p in folder.rglob("*"):
        if file_p.is_file() and file_p.suffix.lower() in supported_exts:
            content = extract_text_from_file(str(file_p))
            if content:
                documents.append(
                    {
                        "id": file_p.name,
                        "filename": file_p.name,
                        "content": content,
                        "path": str(file_p),
                    }
                )

    return documents