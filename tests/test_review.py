"""Tech0 Review のテスト。

実行(リポジトリ直下で): python -m pytest tests/test_review.py
DB はテストごとの一時フォルダに作るので、本番の data/review.db は汚さない。
"""

import io
import sqlite3
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from features.review import db  # noqa: E402
from features.review.checklist_loader import ChecklistError, import_checklist, parse_excel  # noqa: E402


@pytest.fixture
def db_path(tmp_path):
    return tmp_path / "review.db"


def make_excel(path: Path, rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


# ---------------------------------------------------------------- Excel 解析

def test_parse_excel_finds_header_and_skips_rows(tmp_path):
    path = make_excel(tmp_path / "items.xlsx", [
        ["基本設計レビュー項目シート"],               # タイトル行
        [],
        ["No.", "レビュー項目", "観点"],               # 見出しの表記ゆれ(No. / レビュー項目)
        [1, "画面遷移図が記載されている", "主要画面の遷移が図示されているか"],
        [],                                             # 空行
        [2, None, "チェック項目が空の行"],             # 無視される
        [3.0, "エラー処理が定義されている", None],      # No が 3.0、観点なし
    ])

    items, skipped = parse_excel(path)

    assert items == [
        {"item_no": 1, "check_item": "画面遷移図が記載されている", "viewpoint": "主要画面の遷移が図示されているか"},
        {"item_no": 3, "check_item": "エラー処理が定義されている", "viewpoint": ""},
    ]
    assert len(skipped) == 1 and "6 行目" in skipped[0]


def test_parse_excel_without_header_raises(tmp_path):
    path = make_excel(tmp_path / "no_header.xlsx", [["項目A", "説明"], ["x", "y"]])
    with pytest.raises(ChecklistError, match="見出し行"):
        parse_excel(path)


def test_parse_excel_duplicate_no_raises(tmp_path):
    path = make_excel(tmp_path / "dup.xlsx", [
        ["No", "チェック項目", "観点"],
        [1, "項目A", ""],
        [1, "項目B", ""],
    ])
    with pytest.raises(ChecklistError, match="重複"):
        parse_excel(path)


def test_parse_excel_invalid_no_raises(tmp_path):
    path = make_excel(tmp_path / "bad_no.xlsx", [
        ["No", "チェック項目", "観点"],
        ["一", "項目A", ""],
    ])
    with pytest.raises(ChecklistError, match="整数"):
        parse_excel(path)


def test_parse_excel_rejects_non_excel(tmp_path):
    path = tmp_path / "fake.xlsx"
    path.write_text("これは Excel ではない", encoding="utf-8")
    with pytest.raises(ChecklistError, match="読み込めません"):
        parse_excel(path)


# ---------------------------------------------------------------- 登録(置き換え)

def test_import_checklist_replaces_same_phase(db_path):
    first = [{"item_no": i, "check_item": f"項目{i}", "viewpoint": ""} for i in (1, 2, 3)]
    second = [{"item_no": 1, "check_item": "新しい項目", "viewpoint": "新しい観点"}]

    import_checklist("基本設計", first, db_path=db_path)
    import_checklist("要件定義", first, db_path=db_path)
    import_checklist("基本設計", second, db_path=db_path)

    assert db.count_items_by_phase(db_path=db_path) == {"基本設計": 1, "要件定義": 3}
    stored = db.get_review_items("基本設計", db_path=db_path)
    assert stored[0]["check_item"] == "新しい項目"


def test_import_checklist_rejects_unknown_phase(db_path):
    with pytest.raises(ValueError):
        import_checklist("存在しない工程", [{"item_no": 1, "check_item": "x", "viewpoint": ""}], db_path=db_path)


def test_failed_replace_keeps_old_items(db_path):
    """No が重複していて INSERT が失敗しても、元の項目は消えない(トランザクション)。"""
    import_checklist("基本設計", [{"item_no": 1, "check_item": "元の項目", "viewpoint": ""}], db_path=db_path)
    broken = [{"item_no": 1, "check_item": "A", "viewpoint": ""}, {"item_no": 1, "check_item": "B", "viewpoint": ""}]

    with pytest.raises(sqlite3.IntegrityError):
        import_checklist("基本設計", broken, db_path=db_path)

    assert [it["check_item"] for it in db.get_review_items("基本設計", db_path=db_path)] == ["元の項目"]


# ---------------------------------------------------------------- 文書の読み込み(loader)

from docx import Document  # noqa: E402

from features.review import loader  # noqa: E402
from features.review.loader import LoaderError, load_bytes  # noqa: E402


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-8", "cp932"])
def test_decode_text_detects_encoding(encoding):
    text = "要件定義書　第1章 はじめに"
    assert loader._decode_text(text.encode(encoding)) == text


def test_load_txt_in_shift_jis():
    assert load_bytes("memo.txt", "画面遷移図を以下に示す。".encode("cp932")) == "画面遷移図を以下に示す。"


def test_load_csv_joins_cells():
    data = "No,項目\n1,ログイン画面\n\n2,一覧画面\n".encode("utf-8-sig")
    assert load_bytes("list.csv", data) == "No | 項目\n1 | ログイン画面\n2 | 一覧画面"


def test_load_docx_reads_tables_in_order():
    doc = Document()
    doc.add_paragraph("1.2 システム化の方針")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "区分", "方針"
    table.cell(1, 0).text, table.cell(1, 1).text = "基盤方針", "クラウド上に構築する"
    doc.add_paragraph("1.3 設計対象範囲")
    buf = io.BytesIO()
    doc.save(buf)

    assert load_bytes("design.docx", buf.getvalue()) == (
        "1.2 システム化の方針\n区分 | 方針\n基盤方針 | クラウド上に構築する\n1.3 設計対象範囲"
    )


def test_load_normalizes_radicals():
    """PDF に混ざる部首文字(⽬ U+2F6C)は通常の漢字(目)にそろえる。"""
    assert load_bytes("a.txt", "⽬次".encode("utf-8")) == "目次"


def test_load_rejects_unsupported_and_empty():
    with pytest.raises(LoaderError, match="対応していません"):
        load_bytes("image.png", b"\x89PNG")
    with pytest.raises(LoaderError, match="抽出できませんでした"):
        load_bytes("empty.txt", b"  \n ")


# ---------------------------------------------------------------- 評価(engine)。OpenAI はモック

import json  # noqa: E402
from types import SimpleNamespace  # noqa: E402

from features.review import engine  # noqa: E402
from features.review.engine import NO_EVIDENCE, ReviewError, run_review, validate_result  # noqa: E402

BODY = "3.2 画面遷移図を以下に示す。\nログイン画面から一覧画面へ遷移する。"


class FakeClient:
    """chat.completions.create() の代わり。answer(呼ばれた項目のリスト) が返す結果を JSON にして返す。"""

    def __init__(self, answer):
        self.answer = answer
        self.calls: list[list[int]] = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

    def _create(self, **kwargs):
        items = json.loads(kwargs["messages"][1]["content"].split("件)\n", 1)[1].split("\n\n")[0])
        self.calls.append([it["item_no"] for it in items])
        content = json.dumps({"results": self.answer(items)}, ensure_ascii=False)
        message = SimpleNamespace(content=content, refusal=None)
        return SimpleNamespace(choices=[SimpleNamespace(message=message, finish_reason="stop")])


def make_items(n):
    return [{"item_no": i, "check_item": f"項目{i}", "viewpoint": f"観点{i}"} for i in range(1, n + 1)]


def all_ok(items):
    return [{"item_no": it["item_no"], "status": "OK", "evidence": "画面遷移図を以下に示す。", "suggestion": ""} for it in items]


def test_run_review_batches_by_10_and_builds_result():
    client = FakeClient(all_ok)
    result = run_review("基本設計", make_items(25), BODY, "設計書.docx", client=client)

    assert sorted(len(c) for c in client.calls) == [5, 10, 10]  # バッチは同時に投げるので順不同
    assert result["phase"] == "基本設計" and result["file_name"] == "設計書.docx"
    assert result["summary"] == {"ok": 25, "ng": 0}
    assert result["truncated"] is False and result["original_chars"] == len(BODY)
    first = result["results"][0]
    assert first["check_item"] == "項目1" and first["viewpoint"] == "観点1" and first["evidence_found"] is True


def test_run_review_retries_missing_items_once():
    def skip_item_2_first_time(items):
        if len(items) > 1:
            return [r for r in all_ok(items) if r["item_no"] != 2]
        return all_ok(items)

    client = FakeClient(skip_item_2_first_time)
    result = run_review("基本設計", make_items(3), BODY, client=client)
    assert client.calls == [[1, 2, 3], [2]]
    assert [r["status"] for r in result["results"]] == ["OK", "OK", "OK"]


def test_run_review_marks_items_never_returned():
    client = FakeClient(lambda items: [])
    result = run_review("基本設計", make_items(1), BODY, client=client)
    assert result["results"][0]["status"] == "NG"
    assert result["results"][0]["suggestion"] == engine.NOT_JUDGED_SUGGESTION


def test_run_review_truncates_long_text(monkeypatch):
    monkeypatch.setattr(engine, "MAX_DOC_CHARS", 10)
    client = FakeClient(all_ok)
    result = run_review("基本設計", make_items(1), "あ" * 25, client=client)
    assert result["truncated"] is True and result["original_chars"] == 25


def test_validate_result_checks_evidence_and_cleans():
    items = make_items(5)
    raw = [
        {"item_no": 1, "status": "OK", "evidence": "「3.2 画面遷移図を以下に示す。」", "suggestion": "不要な提案"},
        {"item_no": 2, "status": "NG", "evidence": "エラー時は再試行する。", "suggestion": "追記してください。"},
        {"item_no": 3, "status": "NG", "evidence": "(該当する記述なし)", "suggestion": "追記してください。"},
        {"item_no": 4, "status": "OK", "evidence": "ログイン画面から…へ遷移する。", "suggestion": ""},
        {"item_no": 99, "status": "OK", "evidence": "", "suggestion": ""},   # 渡していない番号は捨てる
        {"item_no": 1, "status": "NG", "evidence": "", "suggestion": ""},    # 2 件目は捨てる
    ]
    judged, missing = validate_result(items, raw, BODY)

    assert missing == {5}  # 返ってこなかった項目
    assert judged[1]["evidence_found"] is True and judged[1]["suggestion"] == ""  # かぎかっこ付きでも照合できる
    assert judged[2]["evidence_found"] is False   # 本文にない根拠
    assert judged[3]["evidence_found"] is True    # NG で「該当なし」は照合不要
    assert judged[4]["evidence_found"] is True    # 「…」で省略されていても断片がすべて本文にある
    assert 99 not in judged


def test_ok_without_evidence_is_flagged():
    judged, _ = validate_result(make_items(1), [{"item_no": 1, "status": "OK", "evidence": NO_EVIDENCE, "suggestion": ""}], BODY)
    assert judged[1]["evidence_found"] is False


def test_build_messages_puts_items_before_document():
    messages = engine.build_messages("要件定義", make_items(2), "本文です。")
    assert "「要件定義」工程" in messages[0]["content"]
    user = messages[1]["content"]
    assert user.index("チェック項目") < user.index("<document>") and "本文です。" in user


def test_missing_api_key_raises(monkeypatch):
    import dotenv

    monkeypatch.setattr(dotenv, "load_dotenv", lambda *a, **k: False)  # 本物の .env を読ませない
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    with pytest.raises(ReviewError, match="API キー"):
        run_review("基本設計", make_items(1), BODY)


def test_run_review_without_items_raises():
    with pytest.raises(ReviewError):
        run_review("基本設計", [], BODY, client=FakeClient(all_ok))


def test_run_review_retries_ok_with_unverified_evidence():
    """OK なのに根拠が本文にない回答は、1 回だけ判定し直す。"""
    answers = iter([
        [{"item_no": 1, "status": "OK", "evidence": "本文にない文", "suggestion": ""}],
        [{"item_no": 1, "status": "NG", "evidence": NO_EVIDENCE, "suggestion": "追記してください。"}],
    ])
    client = FakeClient(lambda items: next(answers))
    result = run_review("基本設計", make_items(1), BODY, client=client)
    assert len(client.calls) == 2
    assert result["results"][0]["status"] == "NG" and result["results"][0]["evidence_found"] is True


def test_normalize_text_fixes_pdf_characters():
    assert loader.normalize_text("情報システム部⻑ 受注‧販売 件∕年") == "情報システム部長 受注・販売 件/年"


def test_long_evidence_tolerates_small_gaps():
    """長い引用は、断片の 8 割以上が本文にあれば確認できたとみなす(PDF のページ番号の行が抜けた場合など)。"""
    body = "\n".join(f"第{i}行目の記述はこのとおりである。" for i in range(1, 11)) + "\n第3章 機能要件 11"
    lines = [f"第{i}行目の記述はこのとおりである。" for i in range(1, 11)]
    lines[4] = "AI が言い換えた第5行目の記述。"  # 10 断片中 1 つだけ本文にない
    judged, _ = validate_result(make_items(1), [{"item_no": 1, "status": "OK", "evidence": "\n".join(lines), "suggestion": ""}], body)
    assert judged[1]["evidence_found"] is True

    lines[1] = lines[2] = lines[3] = "本文にない文です。"  # 本文にない断片が 4/10 になると確認できない扱い
    judged, _ = validate_result(make_items(1), [{"item_no": 1, "status": "OK", "evidence": "\n".join(lines), "suggestion": ""}], body)
    assert judged[1]["evidence_found"] is False
