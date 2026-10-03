"""レビュー結果を Excel(.xlsx)にする。ディスクには書かず、バイト列で返す(download_button にそのまま渡す)。

シート「レビュー結果」:
  1〜5 行目 … 工程・対象ファイル・実施日時・判定件数・注意
  7 行目    … 一覧の見出し(No / チェック項目 / 観点 / 判定 / 根拠 / 改善提案 / 備考)
  8 行目〜  … 項目ごとの結果(判定は OK=緑・NG=赤で塗り分け)
シート「NG一覧」: NG の項目だけを同じ形で並べたもの
"""

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from features.review.config import MAX_DOC_CHARS

NOTE = "AI による一次レビューの結果です。最終判断は品質保証部の品質チェック会議で行ってください。"
EVIDENCE_NOT_FOUND = "根拠を本文で確認できませんでした"

COLUMNS = [  # (見出し, 列幅)
    ("No", 6),
    ("チェック項目", 28),
    ("観点", 44),
    ("判定", 8),
    ("根拠（本文からの抜粋）", 48),
    ("改善提案", 48),
    ("備考", 22),
]
TABLE_HEADER_ROW = 7

_HEADER_FILL = PatternFill("solid", fgColor="EAF6FF")
_LABEL_FILL = PatternFill("solid", fgColor="F5F5F5")
_STATUS_STYLE = {
    "OK": (PatternFill("solid", fgColor="E2F4EA"), Font(bold=True, color="23804F")),
    "NG": (PatternFill("solid", fgColor="FBE6E3"), Font(bold=True, color="B93A2E")),
}
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def build_result_xlsx(result: dict) -> bytes:
    """結果データ(engine が作る dict)を Excel のバイト列にする。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "レビュー結果"
    _write_info(ws, result)
    _write_table(ws, result["results"], header_row=TABLE_HEADER_ROW)

    ng_items = [r for r in result["results"] if r["status"] == "NG"]
    _write_table(wb.create_sheet("NG一覧"), ng_items, header_row=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def result_file_name(result: dict) -> str:
    """ダウンロード時のファイル名: レビュー結果_{工程}_{YYYYMMDD_HHMM}.xlsx"""
    stamp = datetime.fromisoformat(result["reviewed_at"]).strftime("%Y%m%d_%H%M")
    phase = re.sub(r'[\\/:*?"<>|]', "_", result["phase"])  # ファイル名に使えない文字を置き換える
    return f"レビュー結果_{phase}_{stamp}.xlsx"


def _note_text(result: dict) -> str:
    if result.get("truncated"):
        return (
            f"{NOTE}\n文書が長いため（{result.get('original_chars', 0):,} 字）、"
            f"先頭 {MAX_DOC_CHARS:,} 字までを評価しました。"
        )
    return NOTE


def _write_info(ws: Worksheet, result: dict) -> None:
    summary = result["summary"]
    reviewed_at = datetime.fromisoformat(result["reviewed_at"]).strftime("%Y/%m/%d %H:%M")
    rows = [
        ("工程", result["phase"]),
        ("対象ファイル", result.get("file_name") or ""),
        ("実施日時", reviewed_at),
        ("判定件数", f"OK {summary['ok']} 件 ／ NG {summary['ng']} 件"),
        ("注意", _note_text(result)),
    ]
    last_col = get_column_letter(len(COLUMNS))
    for i, (label, value) in enumerate(rows, start=1):
        label_cell = ws.cell(row=i, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.fill = _LABEL_FILL
        label_cell.border = _BORDER
        _set_text(ws.cell(row=i, column=2), value).alignment = _WRAP_TOP
        ws.merge_cells(f"B{i}:{last_col}{i}")  # 値は横に長いので B〜最終列を 1 つのセルにする
    ws.row_dimensions[5].height = 32 if result.get("truncated") else 18
    ws.cell(row=4, column=2).font = Font(bold=True)
    ws.cell(row=5, column=2).font = Font(color="B93A2E")


def _write_table(ws: Worksheet, items: list[dict], header_row: int) -> None:
    for col, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    for row, item in enumerate(items, start=header_row + 1):
        values = [
            item["item_no"],
            item["check_item"],
            item.get("viewpoint") or "",
            item["status"],
            item["evidence"],
            item["suggestion"],
            "" if item.get("evidence_found", True) else EVIDENCE_NOT_FOUND,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col)
            if isinstance(value, str):
                _set_text(cell, value)
            else:
                cell.value = value
            cell.alignment = _WRAP_TOP
            cell.border = _BORDER
        status_cell = ws.cell(row=row, column=4)
        if item["status"] in _STATUS_STYLE:
            status_cell.fill, status_cell.font = _STATUS_STYLE[item["status"]]
        status_cell.alignment = Alignment(horizontal="center", vertical="top")

    last_row = header_row + max(len(items), 1)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)  # 見出し行を固定
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{last_row}"


def _set_text(cell, value: str):
    """文字列として書き込む。

    openpyxl は「=」で始まる文字列を数式として保存してしまう。本文や AI の回答に「=」で始まる
    文があっても数式にならないよう、必ず文字列型にする(Excel の数式インジェクション対策)。
    """
    cell.value = value
    cell.data_type = "s"
    return cell
